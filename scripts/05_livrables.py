"""Étape 5 bis - Production des livrables.

  02_Donnees_SIG/   GeoPackage + géodatabase fichier ArcGIS (.gdb) + dictionnaire des données
  05_Tableaux/      classeur Excel des indicateurs (une feuille par thème)
  06_Atlas/         atlas communal PDF (une page par commune, 71 pages + couverture)
  docs/             carte web interactive (GitHub Pages)
"""
import importlib
import json
import re
import shutil
from datetime import date

import geopandas as gpd
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import shapely
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Rectangle
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import cartes as C
import config
from outils import fond_eau, lire, log

etape1 = importlib.import_module("01_consommation_enaf")

DOSSIER_SIG = config.LIVRABLES / "02_Donnees_SIG"
DOSSIER_TAB = config.LIVRABLES / "05_Tableaux"
DOSSIER_ATLAS = config.LIVRABLES / "06_Atlas_communal"

# --- Dictionnaire des données ------------------------------------------------------------
PERIODES = {"1121": "2011-2021", "2125": "2021-2025", "1125": "2011-2025"}
DEST = {"hab": "habitat", "act": "activité économique", "mix": "mixte", "inf": "infrastructures",
        "inc": "destination inconnue"}
CHAMPS = {
    "code_insee": "Code INSEE de la commune",
    "nom": "Nom de la commune",
    "moy_an_1121_ha": "Consommation annuelle moyenne 2011-2021 (ha/an)",
    "moy_an_2125_ha": "Consommation annuelle moyenne 2021-2025 (ha/an)",
    "evol_rythme_pct": "Évolution du rythme annuel 2021-2025 / 2011-2021 (%)",
    "enveloppe_zan_ha": "Enveloppe indicative 2021-2031 = 50 % de la consommation 2011-2021 (ha)",
    "part_enveloppe_pct": "Part de l'enveloppe indicative consommée au 1er janvier 2025 (%)",
    "projection_2131_ha": "Consommation 2021-2031 projetée au rythme 2021-2025 (ha)",
    "surface_com_ha": "Surface de la commune (Cerema, ha)",
    "part_surface_conso_1125_pct": "Part de la surface communale consommée 2011-2025 (%)",
    "pop_2011": "Population 2011 (INSEE)", "pop_2022": "Population 2022 (INSEE)",
    "men_2011": "Ménages 2011 (INSEE)", "men_2022": "Ménages 2022 (INSEE)",
    "emp_2011": "Emplois 2011 (INSEE)", "emp_2022": "Emplois 2022 (INSEE)",
    "m2_hab_par_menage_1122": "m² d'ENAF consommés pour l'habitat par ménage supplémentaire 2011-2022",
    "dominante_1121": "Destination dominante de la consommation 2011-2021 (hab, act, mix, inf, inc)",
    "classe_effort": "Classe d'évolution du rythme de consommation (carte C02)",
    "artif_2019_ha": "Surface artificialisée en 2019 (OCS GE, ha)",
    "artif_2022_ha": "Surface artificialisée en 2022 (OCS GE, ha)",
    "surface_ocs_ha": "Surface couverte par l'OCS GE (ha)",
    "artif_brute_1922_ha": "Artificialisation brute 2019-2022 (ha)",
    "desartif_1922_ha": "Désartificialisation 2019-2022 (ha)",
    "artif_nette_1922_ha": "Artificialisation nette 2019-2022 (ha)",
    "part_artif_2022_pct": "Part du territoire artificialisée en 2022 (%)",
    "zone_U_ha": "Surface en zone U du PLUi (ha)", "zone_AU_ha": "Surface en zone AU (ha)",
    "zone_A_ha": "Surface en zone A (ha)", "zone_N_ha": "Surface en zone N (ha)",
    "au_non_artif_ha": "Surface des zones AU non artificialisée en 2022 (ha)",
    "au_rpg_ha": "Parcelles agricoles RPG 2024 situées en zone AU (ha)",
    "u_non_artif_ha": "Surface non artificialisée en zone U en 2022 (ha)",
    "libelle": "Code de la zone du PLUi", "libelong": "Libellé long de la zone",
    "typezone": "Type de zone CNIG (U, AUc, AUs, A, N)", "famille": "Famille de zone simplifiée",
    "ouverture": "Zone AU ouverte (1AU) ou fermée (2AU)", "vocation": "Vocation de la zone AU",
    "surface_ha": "Surface (ha)", "artif_ha": "Surface artificialisée en 2022 (ha)",
    "non_artif_ha": "Surface non artificialisée en 2022 (ha)",
    "agri_ocs_ha": "Surface agricole OCS GE 2022 (ha)", "rpg_ha": "Surface RPG 2024 (ha)",
    "part_non_artif_pct": "Part non artificialisée (%)", "idurba": "Identifiant du document d'urbanisme",
    "urlfic": "Lien vers le règlement de la zone",
    "type_flux": "Artificialisation ou désartificialisation",
    "cs_2019": "Couverture du sol 2019 (nomenclature OCS GE)", "us_2019": "Usage du sol 2019",
    "cs_2022": "Couverture du sol 2022", "us_2022": "Usage du sol 2022",
    "classe_2019": "Grande classe d'occupation 2019 (1-7)", "classe_2022": "Grande classe d'occupation 2022 (1-7)",
    "commune": "Nom de la commune", "idu": "Identifiant cadastral de la parcelle",
    "zone_plui": "Zone du PLUi (majoritaire)", "type_gisement": "Type de gisement foncier",
    "surface_m2": "Surface de la parcelle (m²)", "surface_libre_m2": "Surface libre constructible (m²)",
    "emprise_bati_m2": "Emprise bâtie (m²)", "coef_emprise": "Coefficient d'emprise au sol",
    "facade_voie_m": "Longueur de contact avec la voie publique (m)",
    "part_boisee": "Part boisée selon l'OCS GE (0-1)", "dans_oap": "Situé dans une OAP (0/1)",
    "vigilance": "Point de vigilance", "logements_bruts": "Potentiel brut de logements",
    "taux_mobilisation": "Taux de mobilisation à 10 ans (hypothèse)",
    "logements_mobilisables": "Logements mobilisables à 10 ans",
    "typepsc": "Type de prescription CNIG", "stypepsc": "Sous-type de prescription",
    "section": "Section cadastrale", "numero": "Numéro de parcelle", "contenance": "Contenance cadastrale (m²)",
    "cleabs": "Identifiant BD TOPO", "nature": "Nature du bâtiment", "usage_1": "Usage principal du bâtiment",
    "construction_legere": "Construction légère (1) ou non (0)", "nombre_de_logements": "Nombre de logements",
    "hauteur": "Hauteur du bâtiment (m)",
}


def decrire(champ):
    if champ in CHAMPS:
        return CHAMPS[champ]
    m = re.fullmatch(r"conso_(\d{4})_ha", champ)
    if m:
        return f"Consommation d'ENAF {PERIODES[m[1]]}, toutes destinations (ha)"
    m = re.fullmatch(r"c(\d{4})_(\w{3})_ha", champ)
    if m:
        return f"Consommation d'ENAF {PERIODES[m[1]]} pour {DEST[m[2]]} (ha)"
    m = re.fullmatch(r"ocs22_classe(\d)_ha", champ)
    if m:
        mod = importlib.import_module("02_evolution_ocsge")
        return f"OCS GE 2022 - {mod.CLASSES[int(m[1])][0]} (ha)"
    return ""


# --- 1. Données SIG -------------------------------------------------------------------------
def en_multipolygone(g):
    """Type géométrique homogène (exigé par la géodatabase ArcGIS) : MultiPolygon."""
    if g is None or g.is_empty:
        return shapely.MultiPolygon()
    if g.geom_type == "MultiPolygon":
        return g
    if g.geom_type == "Polygon":
        return shapely.MultiPolygon([g])
    polys = [p for p in getattr(g, "geoms", []) if p.geom_type in ("Polygon", "MultiPolygon")]
    return shapely.MultiPolygon([q for p in polys for q in getattr(p, "geoms", [p])])


def couches_livrables():
    ind = lire("ind_conso_enaf")
    for nom in ("ind_ocsge", "ind_plui"):
        autre = lire(nom).drop(columns=["geometry", "nom"], errors="ignore")
        ind = ind.merge(autre, on="code_insee", how="left")
    # point intérieur de chaque commune : support des cercles proportionnels (valeurs absolues)
    centroides = ind.copy()
    centroides["geometry"] = ind.geometry.representative_point()
    return {
        "perimetre_metropole": lire("perimetre"),
        "surfaces_eau": fond_eau(),  # Seine et grands plans d'eau (OCS GE 2022), habillage des cartes
        "communes_indicateurs": ind,
        "communes_centroides": centroides,
        "plui_zonage_qualifie": lire("plui_zonage_qualifie").drop(columns=["id_zone"]),
        "plui_zones_au": lire("plui_zones_au").drop(columns=["id_zone"]),
        "ocsge_flux_2019_2022": lire("ocsge_flux_2019_2022"),
        "gisements_densification": lire("gisements_densification"),
        "contraintes_prescriptions": lire("contraintes_prescriptions"),
        # données d'entrée du modèle QGIS de densification (communes test)
        "cadastre_parcelles_etude": lire("cadastre_parcelles"),
        "bdtopo_batiments_etude": batiments_2d(),
    }


def batiments_2d():
    b = lire("bdtopo_batiments")
    b["geometry"] = shapely.force_2d(b.geometry.values)
    return b[["cleabs", "nature", "usage_1", "construction_legere", "nombre_de_logements", "hauteur",
              "code_insee", "geometry"]]


def exporter_sig(c):
    DOSSIER_SIG.mkdir(parents=True, exist_ok=True)
    gpkg = DOSSIER_SIG / "diagnostic_zan_rouen.gpkg"
    gdb = DOSSIER_SIG / "diagnostic_zan_rouen.gdb"
    gpkg.unlink(missing_ok=True)
    shutil.rmtree(gdb, ignore_errors=True)
    for nom, gdf in c.items():
        gdf = gdf.copy()
        for col in gdf.columns:
            if gdf[col].dtype == bool:
                gdf[col] = gdf[col].astype("int16")
        if gdf.geom_type.iloc[0] != "Point":
            gdf = gdf.set_geometry(gdf.geometry.make_valid().apply(en_multipolygone))
            gdf = gdf[~gdf.geometry.is_empty]
        gdf.to_file(gpkg, layer=nom, driver="GPKG", engine="pyogrio")
        gdf.to_file(gdb, layer=nom, driver="OpenFileGDB", engine="pyogrio")
        log(f"  {nom} : {len(gdf)} entités")
    # dictionnaire des données
    lignes = []
    for nom, gdf in c.items():
        for col in gdf.columns:
            if col != "geometry":
                lignes.append({"couche": nom, "champ": col, "type": str(gdf[col].dtype),
                               "description": decrire(col)})
    dico = pd.DataFrame(lignes)
    dico.to_csv(DOSSIER_SIG / "dictionnaire_des_donnees.csv", index=False, encoding="utf-8-sig", sep=";")
    return dico


# --- 2. Classeur Excel ------------------------------------------------------------------------
def classeur(c, dico, syntheses):
    DOSSIER_TAB.mkdir(parents=True, exist_ok=True)
    chemin = DOSSIER_TAB / "Indicateurs_diagnostic_ZAN_Rouen.xlsx"
    ind = pd.DataFrame(c["communes_indicateurs"].drop(columns="geometry"))
    s1, s2, s3, s4 = syntheses
    synth = [
        ("Consommation d'ENAF 2011-2021 (ha)", s1["conso_1121_ha"]),
        ("Consommation d'ENAF 2021-2025 (ha)", s1["conso_2125_ha"]),
        ("Rythme annuel 2011-2021 (ha/an)", s1["moy_an_1121_ha"]),
        ("Rythme annuel 2021-2025 (ha/an)", s1["moy_an_2125_ha"]),
        ("Évolution du rythme (%)", s1["evol_rythme_pct"]),
        ("Enveloppe indicative 2021-2031 (ha)", s1["enveloppe_zan_ha"]),
        ("Part de l'enveloppe consommée au 01/01/2025 (%)", s1["part_enveloppe"] * 100),
        ("Consommation projetée 2021-2031 au rythme actuel (ha)", s1["projection_2031_ha"]),
        ("Surface artificialisée 2022 - OCS GE (ha)", s2["artif_2022_ha"]),
        ("Artificialisation brute 2019-2022 (ha)", s2["artif_brute_ha"]),
        ("Désartificialisation 2019-2022 (ha)", s2["desartif_ha"]),
        ("Artificialisation nette 2019-2022 (ha)", s2["artif_nette_ha"]),
        ("Zones AU du PLUi (ha)", s3["au_total_ha"]),
        ("dont non artificialisées en 2022 (ha)", s3["au_non_artif_ha"]),
        ("Parcelles agricoles RPG 2024 en zone AU (ha)", s3["au_rpg_ha"]),
        ("Enveloppe ZAN restante 2025-2031 (ha)", s3["enveloppe_restante_ha"]),
        ("Surfaces non artificialisées en zone U (ha)", s3["u_non_artif_ha"]),
    ]
    for nom, s in s4.items():
        synth += [(f"Gisements de densification - {nom} (nombre)", s["nb_gisements"]),
                  (f"Gisements de densification - {nom} (ha)", s["surface_ha"]),
                  (f"Logements mobilisables à 10 ans - {nom}", s["logements_mobilisables"])]
    feuilles = {
        "Synthèse": pd.DataFrame(synth, columns=["Indicateur", "Valeur"]),
        "Communes": ind,
        "Série annuelle ENAF": pd.read_csv(config.INTERIM / "conso_annuelle_metropole.csv"),
        "Zones AU": pd.DataFrame(c["plui_zones_au"].drop(columns="geometry")).sort_values("non_artif_ha", ascending=False),
        "Gisements": pd.DataFrame(c["gisements_densification"].drop(columns="geometry")),
        "Transitions OCS GE": pd.read_csv(config.INTERIM / "matrice_transitions_ocsge_2019_2022.csv", index_col=0)
            .reset_index().rename(columns={"index": "Classe 2019 \\ Classe 2022"}),
        "Dictionnaire": dico,
    }
    with pd.ExcelWriter(chemin, engine="openpyxl") as xw:
        for nom, df in feuilles.items():
            df.to_excel(xw, sheet_name=nom, index=False, startrow=2)
            ws = xw.sheets[nom]
            ws["A1"] = f"{config.EPCI_NOM} - Diagnostic foncier ZAN - {nom}"
            ws["A1"].font = Font(bold=True, size=13, color="184F95")
            ws["A2"] = (f"Sources : Cerema 2011-2025, IGN OCS GE 2019-2022, Géoportail de l'Urbanisme, "
                        f"RPG 2024. Réalisation : {config.AUTEUR}, {date.today():%d/%m/%Y}")
            ws["A2"].font = Font(italic=True, size=9, color="52514E")
            entete = PatternFill("solid", fgColor="184F95")
            fin = Side(style="thin", color="E1E0D9")
            for cell in ws[3]:
                cell.fill = entete
                cell.font = Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[3].height = 32
            for j, col in enumerate(df.columns, start=1):
                largeur = min(60, max(10, len(str(col)) + 2,
                                      *(len(f"{v:.1f}" if isinstance(v, float) else str(v)) + 2
                                        for v in df[col].head(200))))
                ws.column_dimensions[get_column_letter(j)].width = largeur
                if pd.api.types.is_float_dtype(df[col]):
                    for row in ws.iter_rows(min_row=4, min_col=j, max_col=j):
                        row[0].number_format = "# ##0.0"
            for row in ws.iter_rows(min_row=4):
                for cell in row:
                    cell.border = Border(bottom=fin)
            ws.freeze_panes = "B4" if nom in ("Communes", "Zones AU", "Gisements") else "A4"
            ws.auto_filter.ref = f"A3:{get_column_letter(len(df.columns))}{len(df) + 3}"
    log(f"  classeur : {chemin.name}")


# --- 3. Atlas communal --------------------------------------------------------------------------
def page_couverture(pdf, s1, n):
    fig = plt.figure(figsize=C.A4_PAYSAGE)
    fig.add_artist(Rectangle((0, 0), 1, 1, transform=fig.transFigure, color=C.ACCENT))
    fig.text(0.07, 0.70, "Atlas communal", color="white", fontsize=40, fontweight="bold")
    fig.text(0.07, 0.62, "Consommation d'espaces, artificialisation et PLUi", color="#dce8f6", fontsize=18)
    fig.text(0.07, 0.56, f"{config.EPCI_NOM} · {n} communes", color="white", fontsize=14)
    fig.text(0.07, 0.52, "Étude indépendante réalisée à partir de données ouvertes - document non officiel",
             color="#dce8f6", fontsize=10)
    fig.text(0.07, 0.30, "Chaque page présente, pour une commune :\n"
             "  - la carte du zonage du PLUi et des sols artificialisés ou renaturés entre 2019 et 2022 ;\n"
             "  - la consommation annuelle d'ENAF 2011-2024 par destination ;\n"
             "  - les indicateurs clés du diagnostic ZAN.", color="white", fontsize=11, linespacing=1.7, va="top")
    fig.text(0.07, 0.07, f"Réalisation : {config.AUTEUR} · {date.today():%B %Y}".replace(
        date.today().strftime("%B"), ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
                                       "septembre", "octobre", "novembre", "décembre"][date.today().month - 1]),
             color="#dce8f6", fontsize=10)
    fig.text(0.07, 0.04, "Sources : Cerema (fichiers fonciers), IGN (OCS GE, ADMIN EXPRESS, BD TOPO), "
             "Géoportail de l'Urbanisme, INSEE", color="#dce8f6", fontsize=8.5)
    pdf.savefig(fig)
    plt.close(fig)


def page_commune(pdf, r, serie, com, zones, flux, gis, eau, num, total):
    fig = plt.figure(figsize=C.A4_PAYSAGE, facecolor="white")
    fig.add_artist(Rectangle((0, 0.915), 1, 0.085, transform=fig.transFigure, color=C.ACCENT))
    fig.text(0.03, 0.958, r.nom, color="white", fontsize=17, fontweight="bold", va="center")
    fig.text(0.97, 0.958, f"INSEE {r.code_insee} · {C.fr(r.pop_2022, 0)} habitants (2022)",
             color="#dce8f6", fontsize=9.5, ha="right", va="center")
    # carte
    ax = fig.add_axes([0.02, 0.07, 0.55, 0.83])
    ax.set_axis_off()
    geom = com[com.code_insee == r.code_insee]
    xmin, ymin, xmax, ymax = geom.total_bounds
    pos = ax.get_position()
    ratio = (pos.height * C.A4_PAYSAGE[1]) / (pos.width * C.A4_PAYSAGE[0])
    cx, cy, w, h = (xmin + xmax) / 2, (ymin + ymax) / 2, (xmax - xmin) * 1.12, (ymax - ymin) * 1.12
    w, h = (h / ratio, h) if h / w > ratio else (w, w * ratio)
    ax.set_xlim(cx - w / 2, cx + w / 2)
    ax.set_ylim(cy - h / 2, cy + h / 2)
    ax.set_aspect("equal")
    ax.add_patch(Rectangle((cx - w / 2, cy - h / 2), w, h, color=C.HORS_TERRITOIRE, zorder=0))
    zc = zones[zones.code_insee == r.code_insee]
    for t, (lib, coul) in C.ZONES.items():
        sel = zc[zc.famille == t]
        if not sel.empty:
            sel.plot(ax=ax, color=coul, alpha=0.30, edgecolor="none", zorder=2, rasterized=True)
    zc[zc.famille == "AU"].plot(ax=ax, facecolor="none", edgecolor=C.ZONES["AU"][1], hatch="////", lw=0.6, zorder=3)
    if eau is not None:
        gpd.clip(eau, geom.buffer(3000)).plot(ax=ax, color=C.EAU, edgecolor="none", zorder=4)
    fc = flux[flux.code_insee == r.code_insee]
    for typ, coul in (("Artificialisation", C.ARTIF), ("Désartificialisation", C.DESARTIF)):
        sel = fc[fc.type_flux == typ]
        if not sel.empty:
            sel.plot(ax=ax, color=coul, edgecolor=coul, lw=0.8, zorder=6)
    gc = gis[gis.code_insee == r.code_insee]
    if not gc.empty:
        gc.plot(ax=ax, color="#1f1f1f", edgecolor="none", zorder=7)
    com.boundary.plot(ax=ax, color="white", lw=1.2, zorder=5)
    geom.boundary.plot(ax=ax, color=C.ENCRE, lw=1.3, zorder=8)
    C.echelle_nord(ax)
    # localisation (colonne de droite, entre les indicateurs et le graphique)
    axl = fig.add_axes([0.60, 0.325, 0.12, 0.14])
    axl.set_axis_off()
    com.plot(ax=axl, color="#e6e4dd", edgecolor="white", lw=0.3, rasterized=True)
    geom.plot(ax=axl, color=C.ACCENT)
    axl.set_aspect("equal")
    axl.set_title("Localisation", fontsize=7.5, color=C.ENCRE_2, loc="left")
    # légende carte
    leg = [(C.ZONES[t][1], C.ZONES[t][0], 0.30) for t in ("U", "AU", "A", "N")] + [
        (C.ARTIF, "Artificialisé 2019-2022", 1), (C.DESARTIF, "Désartificialisé 2019-2022", 1)]
    if not gc.empty:
        leg.append(("#1f1f1f", "Gisement de densification", 1))
    for i, (coul, lib, a) in enumerate(leg):
        y = 0.455 - i * 0.02
        fig.add_artist(Rectangle((0.755, y), 0.014, 0.014, transform=fig.transFigure, facecolor=coul,
                                 alpha=a, edgecolor="none"))
        fig.text(0.774, y + 0.007, lib, fontsize=7, color=C.ENCRE, va="center")
    # indicateurs
    def ligne(y, lib, val, gras=False):
        fig.text(0.60, y, lib, fontsize=8.5, color=C.ENCRE_2, va="center")
        fig.text(0.97, y, val, fontsize=9 if gras else 8.5, color=C.ENCRE, ha="right", va="center",
                 fontweight="bold" if gras else "normal")
        fig.add_artist(plt.Line2D([0.60, 0.97], [y - 0.014] * 2, transform=fig.transFigure, color=C.GRILLE, lw=0.5))
    fig.text(0.60, 0.885, "CONSOMMATION D'ENAF (CEREMA)", fontsize=8.5, fontweight="bold", color=C.ACCENT)
    y = 0.85
    for lib, val, gras in (
        ("2011-2021 (période de référence)", C.fr(r.conso_1121_ha, 1, "ha"), True),
        ("2021-2025", C.fr(r.conso_2125_ha, 1, "ha"), True),
        ("Évolution du rythme annuel", "—" if pd.isna(r.evol_rythme_pct) else C.pct(r.evol_rythme_pct, signe=True), False),
        ("Classe (carte C02)", r.classe_effort.split(" (")[0], False),
        ("Part de l'enveloppe indicative utilisée", "—" if pd.isna(r.part_enveloppe_pct) else C.pct(r.part_enveloppe_pct), False),
        ("m² consommés / ménage supplémentaire", "—" if pd.isna(r.m2_hab_par_menage_1122) else C.fr(r.m2_hab_par_menage_1122, 0, "m²"), False),
    ):
        ligne(y, lib, val, gras)
        y -= 0.030
    y -= 0.012
    fig.text(0.60, y, "OCCUPATION DU SOL (OCS GE) ET PLUi", fontsize=8.5, fontweight="bold", color=C.ACCENT)
    y -= 0.032
    for lib, val in (
        ("Part artificialisée en 2022", C.pct(r.part_artif_2022_pct, 1)),
        ("Artificialisation nette 2019-2022", C.fr(r.artif_nette_1922_ha, 1, "ha")),
        ("Zones AU du PLUi / dont non artificialisées", f"{C.fr(r.zone_AU_ha, 1)} / {C.fr(r.au_non_artif_ha, 1, 'ha')}"),
        ("Terres agricoles (RPG) en zone AU", C.fr(r.au_rpg_ha, 1, "ha")),
    ):
        ligne(y, lib, val)
        y -= 0.030
    # série annuelle
    axb = fig.add_axes([0.62, 0.105, 0.35, 0.165])
    bas = np.zeros(len(serie))
    for k, (lib, coul) in C.DESTINATIONS.items():
        axb.bar(serie.index, serie[k], bottom=bas, color=coul, width=0.7, edgecolor="white", lw=0.5)
        bas += serie[k].to_numpy()
    axb.axvline(2020.5, color=C.AXE, lw=0.6)
    for s in ("top", "right", "left"):
        axb.spines[s].set_visible(False)
    axb.spines["bottom"].set_color(C.AXE)
    axb.yaxis.grid(True, color=C.GRILLE, lw=0.5)
    axb.set_axisbelow(True)
    axb.tick_params(labelsize=7, length=0)
    axb.set_xticks([2011, 2014, 2017, 2020, 2023])
    axb.set_title("Consommation annuelle d'ENAF (ha) par destination", fontsize=8.5, loc="left", color=C.ENCRE)
    x = 0.62
    for lib, coul in C.DESTINATIONS.values():
        lib = lib.split(" (")[0].replace("Destination inconnue", "Inconnue").replace(" habitat / activité", "")
        fig.add_artist(Rectangle((x, 0.058), 0.009, 0.013, transform=fig.transFigure, color=coul))
        fig.text(x + 0.012, 0.0645, lib, fontsize=6.5, color=C.ENCRE_2, va="center")
        x += 0.018 + len(lib) * 0.0046
    fig.text(0.03, 0.025, "Sources : Cerema 2011-2025 ; IGN OCS GE 2019-2022, ADMIN EXPRESS ; "
             "Géoportail de l'Urbanisme (PLUi 04/05/2026) ; ASP RPG 2024 ; INSEE.", fontsize=6.5, color=C.ENCRE_3)
    fig.text(0.97, 0.025, f"{config.AUTEUR} · page {num}/{total}", fontsize=6.5, color=C.ENCRE_3, ha="right")
    pdf.savefig(fig, dpi=200)
    plt.close(fig)


def atlas(c):
    DOSSIER_ATLAS.mkdir(parents=True, exist_ok=True)
    chemin = DOSSIER_ATLAS / "Atlas_communal_ZAN_Metropole_Rouen.pdf"
    ind = c["communes_indicateurs"].sort_values("nom")
    com = lire("communes")
    zones, flux = c["plui_zonage_qualifie"], c["ocsge_flux_2019_2022"]
    gis, eau = c["gisements_densification"], fond_eau()
    d = etape1.charger()
    s1 = json.loads((config.INTERIM / "synthese_etape1.json").read_text(encoding="utf-8"))
    with PdfPages(chemin) as pdf:
        page_couverture(pdf, s1, len(ind))
        for i, (_, r) in enumerate(ind.iterrows(), start=2):
            dc = d[d.idcom == r.code_insee]
            serie = pd.DataFrame({k: [sum(dc[f"art{a % 100:02d}{s}{(a + 1) % 100:02d}"].sum() for s in suff) / 1e4
                                      for a in etape1.ANNEES]
                                  for k, suff in etape1.DEST.items()}, index=etape1.ANNEES)
            page_commune(pdf, r, serie, com, zones, flux, gis, eau, i, len(ind) + 1)
        meta = pdf.infodict()
        meta["Title"] = f"Atlas communal ZAN - {config.EPCI_NOM}"
        meta["Author"] = config.AUTEUR
    log(f"  atlas : {chemin.name} ({len(ind) + 1} pages)")


# --- 4. Carte web --------------------------------------------------------------------------------
def carte_web(c):
    dossier = config.DOCS / "data"
    dossier.mkdir(parents=True, exist_ok=True)
    garder = ["code_insee", "nom", "pop_2022", "conso_1121_ha", "conso_2125_ha", "evol_rythme_pct",
              "classe_effort", "part_enveloppe_pct", "m2_hab_par_menage_1122", "part_artif_2022_pct",
              "artif_nette_1922_ha", "zone_AU_ha", "au_non_artif_ha", "au_rpg_ha", "c1121_hab_ha",
              "c1121_act_ha", "geometry"]
    jeux = {
        "communes": (c["communes_indicateurs"][garder], 15, 2),
        "zones_au": (c["plui_zones_au"][["libelle", "libelong", "vocation", "ouverture", "surface_ha",
                                         "non_artif_ha", "rpg_ha", "urlfic", "code_insee", "geometry"]], 3, 2),
        "flux": (c["ocsge_flux_2019_2022"][["type_flux", "cs_2019", "us_2019", "cs_2022", "us_2022",
                                            "surface_ha", "geometry"]], 1, 3),
        "gisements": (c["gisements_densification"][["commune", "idu", "zone_plui", "type_gisement",
                                                     "surface_libre_m2", "logements_bruts", "vigilance",
                                                     "geometry"]], 0.5, 3),
    }
    lignes = ["// Données générées par scripts/05_livrables.py - ne pas modifier à la main"]
    for nom, (gdf, tol, nd) in jeux.items():
        g = gdf.copy()
        g["geometry"] = g.geometry.simplify(tol, preserve_topology=True)
        g = g.to_crs(4326)
        for col in g.select_dtypes("float").columns:
            g[col] = g[col].round(nd if col != "geometry" else 6)
        g["geometry"] = shapely.set_precision(g.geometry.values, 1e-5)
        js = g.to_json(drop_id=True)
        lignes.append(f"window.DATA_{nom.upper()} = {js};")
    s = {k: json.loads((config.INTERIM / f"synthese_etape{k}.json").read_text(encoding="utf-8")) for k in (1, 2, 3)}
    lignes.append(f"window.SYNTHESE = {json.dumps(s, ensure_ascii=False)};")
    (dossier / "donnees.js").write_text("\n".join(lignes), encoding="utf-8")
    log(f"  carte web : docs/data/donnees.js ({(dossier / 'donnees.js').stat().st_size / 1e6:.1f} Mo)")


def main():
    log("Étape 5 bis - Livrables")
    c = couches_livrables()
    dico = exporter_sig(c)
    syntheses = [json.loads((config.INTERIM / f"synthese_etape{k}.json").read_text(encoding="utf-8")) for k in (1, 2, 3, 4)]
    classeur(c, dico, syntheses)
    carte_web(c)
    atlas(c)
    log("Livrables produits")


if __name__ == "__main__":
    main()
